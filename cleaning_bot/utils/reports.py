from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


TYPE_LABELS = {
    "carpet": "Gilam / Ковёр / Carpet",
    "curtain": "Parda / Шторы / Curtains",
    "bedding": "Adyol-ko'rpa / Одеяло / Blanket",
    "sofa": "Divan qoplamasi / Обивка / Upholstery",
    "other": "Boshqa / Другое / Other",
}


def make_report(rows: list[dict[str, Any]], period: str) -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Completed orders"
    sheet.append(["Order #", "Customer", "Phone", "Address", "Created UTC", "Item type", "Width m", "Length m", "Quantity", "Area per item m2"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="315C4C")
    for row in rows:
        created = datetime.fromisoformat(row["created_at"]).strftime("%Y-%m-%d %H:%M")
        sheet.append([
            row["id"], row["customer_name"], row["phone"], row["address"], created,
            TYPE_LABELS.get(row["item_type"], row["item_type"]), row["width"], row["length"],
            row["quantity"], row["area"],
        ])
    for column in sheet.columns:
        width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 42)
        sheet.column_dimensions[get_column_letter(column[0].column)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
